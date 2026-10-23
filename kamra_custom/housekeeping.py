"""Dirty room report.

Kamra does not store when a room turned dirty, and a room's status is changed
with `frappe.db.set_value`, which skips document hooks. So we do not try to
catch the moment. We work it out from what Kamra already records:

* the room's own `modified` time moves when its status changes, so for a room
  that is Dirty now it is the best "dirty since";
* checkout creates a "Checkout Clean" task linked to the reservation, and the
  reservation carries `actual_check_out`.

No new fields, no hooks, no writes: nothing here can break a Kamra upgrade
except a rename of the fields above, and tests watch those.
"""

from datetime import datetime, timedelta

import frappe
from frappe.utils import get_datetime, now_datetime

from kamra.authz import require_roles

# A Dirty room whose status changed within this long after a checkout is dirty
# because of that checkout. A later change is a manual or room-move event.
CHECKOUT_WINDOW = timedelta(minutes=10)


def build_rows(rooms, modified, tasks, checkouts, now: datetime, last_checkouts=None) -> list[dict]:
	"""Pure function: the report rows, longest wait first.

	rooms      Kamra room_board rows that are Dirty
	modified   {room: datetime} when the room record last changed
	tasks      {room: {"reservation", "status", "assigned_to_user", "name"}}, newest Checkout Clean task
	checkouts  {reservation: datetime} actual_check_out
	last_checkouts  {room: datetime} the room's latest checkout of any stay, whenever the room turned dirty
	"""
	last_checkouts = last_checkouts or {}
	rows = []
	for r in rooms:
		dirty_since = modified.get(r["name"])
		task = tasks.get(r["name"])
		checked_out = checkouts.get(task["reservation"]) if task and task.get("reservation") else None
		from_checkout = bool(
			checked_out and dirty_since and abs(dirty_since - checked_out) <= CHECKOUT_WINDOW
		)
		since = checked_out if from_checkout else dirty_since
		elapsed = int((now - since).total_seconds() // 60) if since else None
		rows.append(
			{
				"room": r["name"],
				"room_number": r["room_number"],
				"floor": r.get("floor"),
				"room_type": r.get("room_type"),
				"occupancy_status": r.get("occupancy_status"),
				"checked_out_at": str(checked_out) if from_checkout else None,
				"last_checkout_at": str(last_checkouts[r["name"]]) if r["name"] in last_checkouts else None,
				"dirty_since": str(dirty_since) if dirty_since else None,
				"source": "Checkout" if from_checkout else "Status change or room move",
				"reason": "Checked out" if from_checkout else "Marked dirty by hand, or moved out of",
				"elapsed_minutes": max(elapsed, 0) if elapsed is not None else None,
				"task_status": task["status"] if task else None,
				"assigned_to": task.get("assigned_to_user") if task else None,
			}
		)
	# longest wait first; rooms with no known time go last
	rows.sort(key=lambda x: (x["elapsed_minutes"] is None, -(x["elapsed_minutes"] or 0), x["room_number"]))
	return rows


@frappe.whitelist()
@require_roles("Housekeeping", "Housekeeping Supervisor", "Front Desk", "Kamra Agent")
def dirty_room_report(property: str) -> dict:
	"""Rooms that are Dirty now, with checkout time and time waiting."""
	from kamra.api import room_board  # Kamra's own board: same roles and property scope

	dirty = [r for r in room_board(property)["rooms"] if r["housekeeping_status"] == "Dirty"]
	names = [r["name"] for r in dirty]
	now = now_datetime()
	if not names:
		return {"property": property, "generated_at": str(now), "rows": []}

	modified = {
		r.name: get_datetime(r.modified)
		for r in frappe.get_all("Room", filters={"name": ("in", names)}, fields=["name", "modified"])
	}

	tasks: dict = {}
	for t in frappe.get_all(
		"Housekeeping Task",
		filters={"room": ("in", names), "task_type": "Checkout Clean"},
		fields=["name", "room", "reservation", "status", "assigned_to_user"],
		order_by="creation desc",
	):
		tasks.setdefault(t.room, t)  # newest first

	reservations = [t.reservation for t in tasks.values() if t.reservation]
	checkouts = {
		r.name: get_datetime(r.actual_check_out)
		for r in frappe.get_all(
			"Reservation",
			filters={"name": ("in", reservations)} if reservations else {"name": ""},
			fields=["name", "actual_check_out"],
		)
		if r.actual_check_out
	}

	last_checkouts = {
		r.room: get_datetime(r.last)
		for r in frappe.db.sql(
			"select room, max(actual_check_out) as last from tabReservation"
			" where room in %(rooms)s and actual_check_out is not null group by room",
			{"rooms": tuple(names)},
			as_dict=True,
		)
		if r.last
	}

	people = {}
	assignees = {t.assigned_to_user for t in tasks.values() if t.assigned_to_user}
	if assignees:
		people = {
			u.name: u.full_name
			for u in frappe.get_all("User", filters={"name": ("in", list(assignees))}, fields=["name", "full_name"])
		}

	rows = build_rows(dirty, modified, tasks, checkouts, now, last_checkouts)
	for row in rows:
		if row["assigned_to"]:
			row["assigned_to"] = people.get(row["assigned_to"]) or row["assigned_to"]
	return {"property": property, "generated_at": str(now), "rows": rows}


XLSX_HEAD = [
	"Room", "Floor", "Room type", "Last checkout date", "Last checkout time", "Dirty since date",
	"Dirty since time", "Waiting (hh:mm)", "Waiting (minutes)", "Reason", "Cleaning task", "Assigned to",
]


def _split(value: str | None) -> tuple[str, str]:
	"""'2026-10-19 13:28:50.4' -> ('2026-10-19', '13:28')."""
	return (value[:10], value[11:16]) if value else ("", "")


def report_sheet_rows(report: dict) -> list[list]:
	"""The downloadable report as rows: a title block, the header, one row per dirty room."""
	rows: list[list] = [
		["Dirty room report"],
		["Property", report["property"]],
		["Generated", report["generated_at"][:16]],
		[],
		XLSX_HEAD,
	]
	for r in report["rows"]:
		cd, ct = _split(r["last_checkout_at"])
		dd, dt = _split(r["dirty_since"])
		m = r["elapsed_minutes"]
		rows.append([
			r["room_number"], r["floor"] or "", (r["room_type"] or "").split("-")[-1],
			cd, ct, dd, dt,
			f"{m // 60}:{m % 60:02d}" if m is not None else "",
			m if m is not None else "",
			r["reason"], r["task_status"] or "", r["assigned_to"] or "",
		])
	return rows


def build_xlsx(report: dict) -> bytes:
	"""An Excel file: real columns, a bold header, widths that fit, header row frozen."""
	from io import BytesIO

	from openpyxl import Workbook
	from openpyxl.styles import Font
	from openpyxl.utils import get_column_letter

	wb = Workbook()
	ws = wb.active
	ws.title = "Dirty rooms"
	rows = report_sheet_rows(report)
	for row in rows:
		ws.append(row)
	ws["A1"].font = Font(bold=True, size=14)
	for cell in ws[5]:
		cell.font = Font(bold=True)
	ws.freeze_panes = "A6"
	for i, name in enumerate(XLSX_HEAD, 1):
		widest = max(len(str(r[i - 1])) for r in rows[4:] if len(r) >= i)
		ws.column_dimensions[get_column_letter(i)].width = min(max(widest, len(name)) + 2, 44)
	out = BytesIO()
	wb.save(out)
	return out.getvalue()


@frappe.whitelist(methods=["GET"])
@require_roles("Housekeeping", "Housekeeping Supervisor", "Front Desk", "Kamra Agent")
def dirty_room_report_xlsx(property: str):
	"""The dirty room report as an Excel file download."""
	report = dirty_room_report(property)
	stamp = report["generated_at"][:16].replace(" ", "-").replace(":", "-")
	frappe.response["filename"] = f"dirty-rooms-{stamp}.xlsx"
	frappe.response["filecontent"] = build_xlsx(report)
	frappe.response["type"] = "binary"
