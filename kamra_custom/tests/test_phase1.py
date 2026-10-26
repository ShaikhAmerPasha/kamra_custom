"""Phase 1: the Excel helper, the feature switches, and the registration card extras."""

import io
import json

import frappe
from frappe.tests import IntegrationTestCase

from kamra_custom import exports, flags, reservations

PROP = "SR Continental"


class TestExports(IntegrationTestCase):
	def test_a_cell_that_could_run_as_a_formula_is_written_as_text(self):
		self.assertEqual(exports.safe_cell("=HYPERLINK(\"x\")"), "'=HYPERLINK(\"x\")")
		for lead in "+-@":
			self.assertTrue(exports.safe_cell(f"{lead}1+1").startswith("'"))
		self.assertEqual(exports.safe_cell("Ravi Kumar"), "Ravi Kumar")
		self.assertEqual(exports.safe_cell(-5), -5)
		self.assertIsNone(exports.safe_cell(None))

	def test_workbook_has_title_header_and_rows(self):
		from openpyxl import load_workbook

		raw = exports.build_xlsx(
			"Plan report", [("guest", "Guest"), ("pax", "Pax")],
			[{"guest": "=cmd", "pax": "2+1"}, {"guest": "Meera"}], ("Property: SR Continental",),
		)
		ws = load_workbook(io.BytesIO(raw)).active
		rows = [[c for c in r] for r in ws.iter_rows(values_only=True)]
		self.assertEqual(rows[0][0], "Plan report")
		self.assertEqual(rows[1][0], "Property: SR Continental")
		self.assertEqual(rows[3], ["Guest", "Pax"])
		self.assertEqual(rows[4], ["'=cmd", "2+1"])
		self.assertEqual(rows[5], ["Meera", None])

	def test_too_many_rows_are_refused(self):
		with self.assertRaises(frappe.ValidationError):
			exports.build_xlsx("x", [("a", "A")], [{"a": 1}] * (exports.MAX_ROWS + 1))

	def test_response_is_a_binary_download(self):
		exports.send_xlsx("report", b"abc")
		self.assertEqual(frappe.response["filename"], "report.xlsx")
		self.assertEqual(frappe.response["type"], "binary")


class TestFlags(IntegrationTestCase):
	def tearDown(self):
		frappe.db.rollback()

	def test_a_switch_never_saved_keeps_its_default(self):
		out = flags.ui_flags()
		self.assertEqual(set(out), set(flags.FLAGS))
		for name, default in flags.DEFAULTS.items():
			frappe.db.sql("delete from tabSingles where doctype = 'Kamra Custom Settings' and field = %s", name)
		self.assertEqual(flags.ui_flags(), flags.DEFAULTS)
		# the bill layout needs the accountant's word, so it starts off
		self.assertFalse(flags.DEFAULTS["merge_meal_in_room_line"])

	def test_a_switch_can_be_turned_off(self):
		frappe.db.set_single_value("Kamra Custom Settings", "grc_staff_name", 0)
		self.assertFalse(flags.ui_flags()["grc_staff_name"])
		self.assertTrue(flags.ui_flags()["checkout_date_on_booking"])


class TestGrcExtra(IntegrationTestCase):
	def setUp(self):
		frappe.set_user("Administrator")
		self.name = frappe.db.get_value("Reservation", {"property": PROP})
		if not self.name:
			self.skipTest("no reservation on this site")
		self.doc = frappe.get_doc("Reservation", self.name)

	def tearDown(self):
		frappe.db.rollback()
		frappe.set_user("Administrator")

	def set_steps(self, **card):
		frappe.db.set_value("Reservation", self.name, "checkin_steps", json.dumps({"registration_card": card}))

	def test_staff_name_is_the_user_who_checked_in(self):
		frappe.db.set_value("Reservation", self.name, "checked_in_by", "Administrator")
		self.assertEqual(reservations.grc_extra(self.name)["staff_name"], "Administrator")

	def test_before_check_in_it_is_who_took_the_card_step_then_the_viewer(self):
		frappe.db.set_value("Reservation", self.name, "checked_in_by", None)
		self.set_steps(state="done", by="Administrator", method="paper")
		self.assertEqual(reservations.grc_extra(self.name)["staff_name"], "Administrator")
		frappe.db.set_value("Reservation", self.name, "checkin_steps", None)
		self.assertTrue(reservations.grc_extra(self.name)["staff_name"])

	def test_paper_signature_is_reported_only_when_the_step_is_done_on_paper(self):
		self.set_steps(state="done", by="Administrator", method="paper")
		self.assertTrue(reservations.grc_extra(self.name)["signed_on_paper"])
		self.set_steps(state="done", by="Administrator", method="screen")
		self.assertFalse(reservations.grc_extra(self.name)["signed_on_paper"])
		self.set_steps(state="skipped", by="Administrator", method="paper")
		self.assertFalse(reservations.grc_extra(self.name)["signed_on_paper"])

	def test_bad_step_data_does_not_break_the_card(self):
		frappe.db.set_value("Reservation", self.name, "checkin_steps", "not json")
		self.assertFalse(reservations.grc_extra(self.name)["signed_on_paper"])


class TestKamraContract(IntegrationTestCase):
	"""Names we read from Kamra. If a Kamra update renames one, this fails here, not at the desk."""

	def test_reservation_fields_we_read(self):
		meta = frappe.get_meta("Reservation")
		for field in ("checked_in_by", "checkin_steps", "precheckin_signature", "precheckin_consent", "room_type",
		              "meal_plan", "status", "cancellation_reason", "cancellation_number", "cancellation_fee",
		              "cancelled_on"):
			self.assertTrue(meta.has_field(field), field)

	def test_checkin_step_record_shape(self):
		import inspect

		from kamra import checkin

		self.assertTrue({"signature", "on_paper"} <= set(inspect.signature(checkin.sign_registration_card).parameters))
		src = inspect.getsource(checkin.sign_registration_card)
		self.assertIn('"paper"', src)
		self.assertIn('"registration_card"', src)
