"""Front desk reports: the right people, the filters, Excel, and nothing off-limits."""

import io

import frappe
from frappe.tests import IntegrationTestCase
from frappe.utils import add_days, nowdate

from kamra_custom import desk_reports, overrides

PROP = "SR Continental"
DLX = "SR Continental-DLX"
TODAY = nowdate()
TOMORROW = add_days(TODAY, 1)


class TestDeskReports(IntegrationTestCase):
	def setUp(self):
		frappe.set_user("Administrator")
		if not frappe.db.exists("Room Type", DLX):
			self.skipTest("demo property not on this site")
		# these tests assume Kamra's own way: a room is given at booking. The hotel's switches must not change that.
		frappe.db.set_single_value("Kamra Custom Settings", "assign_room_at_checkin", 0)
		frappe.db.set_single_value("Kamra Custom Settings", "clean_rooms_only", 0)
		self.cp = frappe.db.get_value("Meal Plan", {"property": PROP, "code": "CP"})

	def tearDown(self):
		frappe.db.rollback()
		frappe.set_user("Administrator")

	def book(self, name, phone, **kw):
		args = dict(property=PROP, room_type=DLX, check_in_date=TODAY, check_out_date=TOMORROW,
		            guest_name=name, phone=phone, adults=2, children=1)
		args.update(kw)
		return overrides.create_booking(**args)["reservation"]

	def run_report(self, key, **kw):
		return desk_reports.run(key, PROP, kw.pop("from_date", TODAY), kw.pop("to_date", TODAY), **kw)

	def refs(self, out):
		return {r["ref"] for r in out["rows"]}

	def test_arrivals_include_those_already_checked_in_with_a_status(self):
		a = self.book("Arr One", "9000066601")
		b = self.book("Arr Two", "9000066602")
		from kamra.api import check_in
		check_in(b)
		rows = {r["ref"]: r for r in self.run_report("arrivals")["rows"]}
		self.assertEqual(rows[a]["status"], "Expected")
		self.assertEqual(rows[b]["status"], "Arrived")

	def test_the_date_is_a_filter_not_fixed(self):
		later = self.book("Arr Later", "9000066603", check_in_date=add_days(TODAY, 5), check_out_date=add_days(TODAY, 6))
		self.assertNotIn(later, self.refs(self.run_report("arrivals")))
		out = self.run_report("arrivals", from_date=add_days(TODAY, 5), to_date=add_days(TODAY, 5))
		self.assertIn(later, self.refs(out))

	def test_packs_are_adults_plus_children(self):
		ref = self.book("Pack Test", "9000066604")
		row = next(r for r in self.run_report("arrivals")["rows"] if r["ref"] == ref)
		self.assertEqual((row["adults"], row["children"], row["packs"]), (2, 1, 3))

	def test_departures_list_who_is_due_out(self):
		ref = self.book("Dep Test", "9000066605")
		from kamra.api import check_in
		check_in(ref)
		row = next(r for r in self.run_report("departures", from_date=TOMORROW, to_date=TOMORROW)["rows"] if r["ref"] == ref)
		self.assertEqual(row["status"], "Due out")

	def test_in_house_shows_guest_room_packs_plan_and_dates(self):
		ref = self.book("House Test", "9000066606", meal_plan=self.cp)
		from kamra.api import check_in
		check_in(ref)
		out = self.run_report("in_house")
		row = next(r for r in out["rows"] if r["ref"] == ref)
		self.assertEqual((row["guest"], row["packs"], row["check_in"], row["check_out"]), ("House Test", 3, getdate_(TODAY), getdate_(TOMORROW)))
		self.assertTrue(row["room"])
		self.assertEqual(row["plan"], "CP")
		self.assertTrue(out["summary"])

	def test_the_plan_filter_narrows_the_list(self):
		a = self.book("Plan A", "9000066607", meal_plan=self.cp)
		b = self.book("Plan B", "9000066608")
		from kamra.api import check_in
		check_in(a)
		check_in(b)
		self.assertEqual(self.refs(self.run_report("plan", plan="CP")) & {a, b}, {a})
		self.assertEqual(self.refs(self.run_report("plan", plan="Room only")) & {a, b}, {b})

	def test_a_guest_who_has_not_arrived_is_not_in_house(self):
		ref = self.book("Not Yet", "9000066609")
		self.assertNotIn(ref, self.refs(self.run_report("in_house")))

	def test_agreed_rates_by_approver(self):
		a = self.book("Agree MD", "9000066610", custom_rate={"rates": [2300], "approver": "MD Sir"})
		b = self.book("Agree GM", "9000066611", custom_rate={"rates": [2300], "approver": "GM Sir"})
		both = self.run_report("agreed_rates")
		self.assertTrue({a, b} <= self.refs(both))
		only = self.run_report("agreed_rates", approver="MD Sir")
		self.assertEqual(self.refs(only) & {a, b}, {a})
		self.assertTrue(all(r["approver"] == "MD Sir" for r in only["rows"]))

	def test_an_unassigned_room_reads_not_assigned_yet(self):
		frappe.db.set_single_value("Kamra Custom Settings", "assign_room_at_checkin", 1)
		ref = self.book("No Room", "9000066612")
		row = next(r for r in self.run_report("arrivals")["rows"] if r["ref"] == ref)
		self.assertEqual(row["room"], "Not assigned yet")

	def test_excel_holds_the_same_rows_and_is_formula_safe(self):
		from openpyxl import load_workbook

		self.book("=HYPERLINK(\"x\")", "9000066613")
		desk_reports.download("arrivals", PROP, TODAY, TODAY)
		content = frappe.response["filecontent"]
		self.assertTrue(frappe.response["filename"].endswith(".xlsx"))
		sheet = load_workbook(io.BytesIO(content)).active
		cells = [c.value for row in sheet.iter_rows() for c in row if isinstance(c.value, str)]
		self.assertIn("'=HYPERLINK(\"x\")", cells)
		self.assertNotIn("=HYPERLINK(\"x\")", cells)

	def test_bad_input_is_refused(self):
		with self.assertRaises(frappe.ValidationError):
			desk_reports.run("nope", PROP)
		with self.assertRaises(frappe.ValidationError):
			desk_reports.run("arrivals", PROP, TOMORROW, TODAY)
		with self.assertRaises(frappe.ValidationError):
			desk_reports.run("arrivals", PROP, TODAY, add_days(TODAY, 900))

	def test_only_desk_roles_may_read(self):
		frappe.set_user("Guest")
		with self.assertRaises(frappe.PermissionError):
			desk_reports.run("arrivals", PROP)


def getdate_(s):
	from frappe.utils import getdate

	return getdate(s)


class TestCatalog(IntegrationTestCase):
	def test_lists_every_report_and_the_filter_choices(self):
		frappe.set_user("Administrator")
		out = desk_reports.catalog(PROP)
		self.assertEqual({r["key"] for r in out["reports"]}, set(desk_reports.REPORTS))
		self.assertEqual(out["plans"][0], "Room only")
		self.assertIn("MD Sir", out["approvers"])


class TestCancelAndAmend(TestDeskReports):
	def test_a_cancelled_booking_shows_who_why_and_the_number(self):
		from kamra.api import cancel_reservation

		ref = self.book("Cancel Me", "9000066620", check_in_date=add_days(TODAY, 30), check_out_date=add_days(TODAY, 31))
		cancel_reservation(ref, reason="Change of plan", note="Called the desk")
		row = next(r for r in self.run_report("cancellations")["rows"] if r["ref"] == ref)
		self.assertTrue(row["reason"])
		self.assertEqual(row["note"], "Called the desk")
		self.assertTrue(row["cancel_no"])
		self.assertTrue(row["cancelled_by"])
		self.assertNotIn(ref, self.refs(self.run_report("cancellations", from_date=add_days(TODAY, 3), to_date=add_days(TODAY, 4))))

	def test_an_amended_stay_shows_the_new_dates_and_who_did_it(self):
		from kamra.api import amend_stay

		ref = self.book("Amend Me", "9000066621", check_in_date=add_days(TODAY, 30), check_out_date=add_days(TODAY, 31))
		new_out = add_days(TODAY, 32)
		amend_stay(ref, add_days(TODAY, 30), new_out, confirm_price=1)
		rows = [r for r in self.run_report("amendments")["rows"] if r["ref"] == ref]
		row = next(r for r in rows if r["what"] == "Dates changed")
		self.assertIn(new_out, row["details"])
		self.assertTrue(row["changed_by"])

	def test_a_room_chosen_in_the_check_in_steps_is_listed(self):
		frappe.db.set_single_value("Kamra Custom Settings", "assign_room_at_checkin", 1)
		ref = self.book("Assign Me", "9000066622")
		from kamra.checkin import assign_checkin_room
		room = frappe.db.get_value("Room", {"room_type": DLX, "housekeeping_status": "Clean", "occupancy_status": "Vacant"})
		assign_checkin_room(ref, room=room)
		rows = [r for r in self.run_report("amendments")["rows"] if r["ref"] == ref]
		self.assertTrue(any(room in r["details"] for r in rows), rows)


class TestFindBooking(TestDeskReports):
	def test_finds_by_name_number_and_booking_ref(self):
		from kamra_custom import reservations

		ref = self.book("Findable Person", "9000066630")
		for q in ("Findable", "66630", ref):
			self.assertIn(ref, [r["name"] for r in reservations.find_to_act_on("cancel", PROP, q)], q)
		self.assertNotIn(ref, [r["name"] for r in reservations.find_to_act_on("cancel", PROP, "Nobody Here")])

	def test_cancel_lists_confirmed_only_but_amend_lists_checked_in_too(self):
		from kamra.api import check_in
		from kamra_custom import reservations

		ref = self.book("Stay Person", "9000066631")
		check_in(ref)
		self.assertNotIn(ref, [r["name"] for r in reservations.find_to_act_on("cancel", PROP, "Stay Person")])
		self.assertIn(ref, [r["name"] for r in reservations.find_to_act_on("amend", PROP, "Stay Person")])

	def test_a_blank_search_lists_next_arrivals_and_a_bad_action_is_refused(self):
		from kamra_custom import reservations

		ref = self.book("Soon Person", "9000066632")
		self.assertIn(ref, [r["name"] for r in reservations.find_to_act_on("cancel", PROP)])
		with self.assertRaises(frappe.ValidationError):
			reservations.find_to_act_on("delete", PROP)


class TestComplimentaryReport(TestDeskReports):
	def test_lists_what_was_given_and_by_whom(self):
		from kamra_custom import complimentary

		ref = self.book("Comp Guest", "9000066640", meal_plan=self.cp)
		worth = frappe.db.get_value("Reservation", ref, "amount_after_tax")
		other = self.book("Paying Guest", "9000066641")
		complimentary.mark(ref, "GM Sir", "Owner's friend")
		out = self.run_report("complimentary")
		self.assertEqual(self.refs(out) & {ref, other}, {ref})
		row = next(r for r in out["rows"] if r["ref"] == ref)
		self.assertEqual((row["approver"], row["reason"]), ("GM Sir", "Owner's friend"))
		self.assertAlmostEqual(row["value"], worth, places=1)
		self.assertEqual(self.refs(self.run_report("complimentary", approver="MD Sir")) & {ref}, set())
		self.assertTrue(out["summary"])


class TestAmendmentsNotRepeated(TestDeskReports):
	def test_a_date_change_is_one_row_not_three(self):
		from kamra.api import amend_stay

		ref = self.book("Once Only", "9000066650", check_in_date=add_days(TODAY, 30), check_out_date=add_days(TODAY, 31))
		amend_stay(ref, add_days(TODAY, 30), add_days(TODAY, 33), confirm_price=1)
		rows = [r for r in self.run_report("amendments")["rows"] if r["ref"] == ref]
		self.assertEqual([r["what"] for r in rows], ["Dates changed"])
