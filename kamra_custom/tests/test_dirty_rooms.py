"""Phase 3: the dirty room report."""

from datetime import datetime, timedelta

import frappe
from frappe.tests import IntegrationTestCase

from kamra_custom.housekeeping import CHECKOUT_WINDOW, build_rows, build_xlsx, report_sheet_rows

NOW = datetime(2026, 10, 17, 14, 0, 0)


def room(name, number, floor="1"):
	return {"name": name, "room_number": number, "floor": floor, "room_type": "DLX", "occupancy_status": "Vacant"}


class TestBuildRows(IntegrationTestCase):
	def test_checkout_dirty_room_reports_checkout_time_and_wait(self):
		out = NOW - timedelta(hours=2, minutes=10)
		rows = build_rows(
			[room("P-101", "101")],
			{"P-101": out + timedelta(seconds=2)},
			{"P-101": {"reservation": "R1", "status": "Pending", "assigned_to_user": None}},
			{"R1": out},
			NOW,
		)
		self.assertEqual(rows[0]["source"], "Checkout")
		self.assertEqual(rows[0]["checked_out_at"], str(out))
		self.assertEqual(rows[0]["elapsed_minutes"], 130)

	def test_status_changed_long_after_checkout_is_not_a_checkout(self):
		# checked out yesterday, cleaned, then marked dirty again this morning
		out = NOW - timedelta(days=1)
		rows = build_rows(
			[room("P-101", "101")],
			{"P-101": NOW - timedelta(hours=3)},
			{"P-101": {"reservation": "R1", "status": "Done", "assigned_to_user": None}},
			{"R1": out},
			NOW,
		)
		self.assertEqual(rows[0]["source"], "Status change or room move")
		self.assertIsNone(rows[0]["checked_out_at"])
		self.assertEqual(rows[0]["elapsed_minutes"], 180)

	def test_room_with_no_task_uses_status_change_time(self):
		rows = build_rows([room("P-102", "102")], {"P-102": NOW - timedelta(minutes=45)}, {}, {}, NOW)
		self.assertEqual(rows[0]["elapsed_minutes"], 45)
		self.assertIsNone(rows[0]["task_status"])

	def test_longest_wait_first_unknown_last(self):
		rows = build_rows(
			[room("P-1", "1"), room("P-2", "2"), room("P-3", "3")],
			{"P-1": NOW - timedelta(minutes=5), "P-3": NOW - timedelta(minutes=90)},
			{},
			{},
			NOW,
		)
		self.assertEqual([r["room_number"] for r in rows], ["3", "1", "2"])

	def test_wait_never_negative(self):
		rows = build_rows([room("P-1", "1")], {"P-1": NOW + timedelta(minutes=3)}, {}, {}, NOW)
		self.assertEqual(rows[0]["elapsed_minutes"], 0)

	def test_window_boundary(self):
		out = NOW - timedelta(hours=1)
		inside = build_rows(
			[room("P-1", "1")], {"P-1": out + CHECKOUT_WINDOW},
			{"P-1": {"reservation": "R", "status": "Pending"}}, {"R": out}, NOW,
		)
		outside = build_rows(
			[room("P-1", "1")], {"P-1": out + CHECKOUT_WINDOW + timedelta(seconds=1)},
			{"P-1": {"reservation": "R", "status": "Pending"}}, {"R": out}, NOW,
		)
		self.assertEqual(inside[0]["source"], "Checkout")
		self.assertNotEqual(outside[0]["source"], "Checkout")


class TestEndpoint(IntegrationTestCase):
	"""Against the site's own data. Needs a property with at least one room."""

	def setUp(self):
		frappe.set_user("Administrator")
		self.room = frappe.db.get_value("Room", {}, ["name", "property", "housekeeping_status"], as_dict=True)
		if not self.room:
			self.skipTest("no room on this site")

	def test_dirty_room_appears_with_a_time(self):
		from kamra_custom.housekeeping import dirty_room_report

		frappe.db.set_value("Room", self.room.name, "housekeeping_status", "Dirty")
		out = dirty_room_report(self.room.property)
		mine = [r for r in out["rows"] if r["room"] == self.room.name]
		self.assertEqual(len(mine), 1)
		self.assertIsNotNone(mine[0]["dirty_since"])
		self.assertGreaterEqual(mine[0]["elapsed_minutes"], 0)
		self.assertIn("generated_at", out)

	def test_clean_room_is_not_listed(self):
		from kamra_custom.housekeeping import dirty_room_report

		frappe.db.set_value("Room", self.room.name, "housekeeping_status", "Clean")
		out = dirty_room_report(self.room.property)
		self.assertNotIn(self.room.name, [r["room"] for r in out["rows"]])

	def test_kamra_fields_we_read_still_exist(self):
		# a Kamra rename must fail here, not in the housekeeping desk
		task = frappe.get_meta("Housekeeping Task")
		for f in ("room", "reservation", "task_type", "status", "assigned_to_user"):
			self.assertTrue(task.has_field(f), f)
		self.assertTrue(frappe.get_meta("Reservation").has_field("actual_check_out"))
		self.assertTrue(frappe.get_meta("Room").has_field("housekeeping_status"))


class TestLastCheckout(IntegrationTestCase):
	def test_a_room_marked_dirty_by_hand_still_shows_its_last_checkout(self):
		last = NOW - timedelta(days=3)
		rows = build_rows(
			[room("P-103", "103")], {"P-103": NOW - timedelta(minutes=20)}, {}, {}, NOW,
			last_checkouts={"P-103": last},
		)
		self.assertEqual(rows[0]["last_checkout_at"], str(last))
		self.assertIsNone(rows[0]["checked_out_at"])
		self.assertEqual(rows[0]["elapsed_minutes"], 20)

	def test_no_last_checkout_is_none(self):
		rows = build_rows([room("P-1", "1")], {"P-1": NOW}, {}, {}, NOW)
		self.assertIsNone(rows[0]["last_checkout_at"])


class TestExcelReport(IntegrationTestCase):
	REPORT = {
		"property": "SR Continental", "generated_at": "2026-10-19 14:04:11.5",
		"rows": [
			{"room": "SR-201", "room_number": "201", "floor": "2", "room_type": "SR-PRM",
			 "checked_out_at": "2026-10-19 11:40:00", "last_checkout_at": "2026-10-19 11:40:00",
			 "dirty_since": "2026-10-19 11:40:01", "source": "Checkout", "reason": "Checked out",
			 "elapsed_minutes": 144, "task_status": "Pending", "assigned_to": "Asha K"},
			{"room": "SR-103", "room_number": "103", "floor": None, "room_type": None,
			 "checked_out_at": None, "last_checkout_at": None, "dirty_since": "2026-10-19 13:40:19",
			 "source": "x", "reason": "Marked dirty by hand, or moved out of",
			 "elapsed_minutes": None, "task_status": None, "assigned_to": None},
		],
	}

	def test_each_detail_is_its_own_cell(self):
		rows = report_sheet_rows(self.REPORT)
		self.assertEqual(rows[1], ["Property", "SR Continental"])
		self.assertEqual(rows[2], ["Generated", "2026-10-19 14:04"])
		self.assertEqual(rows[4][:5], ["Room", "Floor", "Room type", "Last checkout date", "Last checkout time"])
		self.assertEqual(rows[5][:9], ["201", "2", "PRM", "2026-10-19", "11:40", "2026-10-19", "11:40", "2:24", 144])
		self.assertEqual(rows[5][9:], ["Checked out", "Pending", "Asha K"])

	def test_unknown_values_are_blank_cells(self):
		row = report_sheet_rows(self.REPORT)[6]
		self.assertEqual(row[:9], ["103", "", "", "", "", "2026-10-19", "13:40", "", ""])

	def test_the_file_is_an_excel_workbook_with_columns(self):
		from io import BytesIO

		from openpyxl import load_workbook

		ws = load_workbook(BytesIO(build_xlsx(self.REPORT))).active
		self.assertEqual(ws["A1"].value, "Dirty room report")
		self.assertEqual([c.value for c in ws[5]][:3], ["Room", "Floor", "Room type"])
		self.assertEqual(ws["D6"].value, "2026-10-19")
		self.assertEqual(ws["E6"].value, "11:40")
		self.assertEqual(ws.freeze_panes, "A6")
		self.assertTrue(ws["A5"].font.bold)

	def test_the_download_endpoint_returns_a_file(self):
		frappe.set_user("Administrator")
		prop = frappe.db.get_value("Property", {}, "name")
		if not prop:
			self.skipTest("no property")
		from kamra_custom.housekeeping import dirty_room_report_xlsx

		dirty_room_report_xlsx(prop)
		self.assertTrue(frappe.response["filename"].endswith(".xlsx"))
		self.assertEqual(frappe.response["type"], "binary")
		self.assertTrue(frappe.response["filecontent"].startswith(b"PK"))


class TestRealCheckout(IntegrationTestCase):
	"""A guest checks out through Kamra; the report shows when, in the sheet too."""

	def test_checkout_time_reaches_the_report_and_the_excel_file(self):
		frappe.set_user("Administrator")
		from kamra.api import add_folio_payment, check_out

		res = None
		for cand in frappe.get_all(
			"Reservation", filters={"status": "Checked In", "room": ("is", "set")}, fields=["name", "room", "property"]
		):
			frappe.db.savepoint("try_checkout")
			try:
				for f in frappe.get_all("Folio", filters={"reservation": cand.name, "status": "Open"}, fields=["name", "balance"]):
					if f.balance and f.balance > 0:
						add_folio_payment(f.name, "Cash", f.balance)
				check_out(cand.name)
				res = cand
				break
			except frappe.ValidationError:
				frappe.db.rollback(save_point="try_checkout")  # this stay cannot leave yet; try the next
		if not res:
			self.skipTest("no guest in house that can check out on this site")

		from kamra_custom.housekeeping import dirty_room_report

		report = dirty_room_report(res.property)
		row = next(r for r in report["rows"] if r["room"] == res.room)
		self.assertEqual(row["reason"], "Checked out")
		self.assertTrue(row["checked_out_at"])
		self.assertEqual(row["checked_out_at"], row["last_checkout_at"])
		sheet = report_sheet_rows(report)
		mine = next(r for r in sheet[5:] if r[0] == res.room.split("-")[-1])
		self.assertEqual(mine[3], row["checked_out_at"][:10])
		self.assertEqual(mine[4], row["checked_out_at"][11:16])
		frappe.db.rollback()
